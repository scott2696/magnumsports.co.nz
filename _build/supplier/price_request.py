#!/usr/bin/env python3
"""Price requests for suppliers that quote rather than publish prices.

    python3 _build/supplier/price_request.py make ultra-safety
        writes "~/Documents/<Supplier> price request - Magnum Sports.xlsx": every
        product we would sell (no flag, not excluded) that has no price yet,
        with yellow cells for the supplier's unit price and minimum order.

    python3 _build/supplier/price_request.py load ultra-safety <file.xlsx>
        reads the filled-in prices back into the supplier's sheet
        (supplier_fob_usd, supplier_moq), matched on our SKU. Then run
        reprice.py and fetch_images.py, check the new photos, and rebuild.
Needs openpyxl.
"""
import csv, json, os, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
SUPPLIERS = json.load(open(os.path.join(HERE, "suppliers.json"), encoding="utf-8"))
HARMLESS = {"", "Department is a guess - check"}
EXCLUDE = {l.split("#")[0].strip() for l in open(os.path.join(HERE, "exclude.txt"))} - {""}
HDR = ["#", "Our SKU", "Supplier item / model", "Product (supplier title)", "Department",
       "Supplier page", "Unit price (USD)", "Minimum order"]


def rows_for(key):
    info = SUPPLIERS[key]
    path = os.path.join(HERE, info["sheet"])
    return info, path, list(csv.DictReader(open(path, encoding="utf-8-sig")))


def make(key):
    info, _, rows = rows_for(key)
    want = [r for r in rows if r["flags"].strip() in HARMLESS and r["sku"] not in EXCLUDE
            and not r["supplier_fob_usd"].strip()]
    want.sort(key=lambda r: (r["dept"], r["supplier_title"].lower()))
    F = lambda **k: Font(name="Arial", **k)
    wb = Workbook()
    ws = wb.active
    ws.title = "Price request"
    ws["A1"] = f"Magnum Sports: price request for {info['company']}"
    ws["A1"].font = F(bold=True, size=14)
    ws["A2"] = (f"{len(want)} products from {info['website']} we would like to stock. Please enter your unit price "
                "in US dollars and your minimum order quantity in the yellow cells, e.g. 12.50 and 50.")
    ws["A3"] = "Keep the 'Our SKU' column unchanged: we use it to match your prices to our catalogue."
    for c in ("A2", "A3"):
        ws[c].font = F(size=9, italic=True, color="595959")
    H = 5
    for j, h in enumerate(HDR, 1):
        c = ws.cell(row=H, column=j, value=h)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B2836")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    yellow = PatternFill("solid", fgColor="FFFF00")
    for n, r in enumerate(want, 1):
        i = H + n
        vals = [n, r["sku"], r["model"] or "—", r["supplier_title"], r["dept"], "Open", None, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F(size=10)
            c.alignment = Alignment(vertical="top", wrap_text=(j == 4))
        link = ws.cell(row=i, column=6)
        link.hyperlink = r["supplier_url"]
        link.font = F(size=10, color="0563C1", underline="single")
        ws.cell(row=i, column=7).fill = yellow
        ws.cell(row=i, column=7).number_format = '"$"#,##0.00'
        ws.cell(row=i, column=8).fill = yellow
    for col, w in zip("ABCDEFGH", [5, 26, 18, 70, 20, 10, 14, 13]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=H + 1, column=5)
    out = os.path.expanduser(f"~/Documents/{info['name']} price request - Magnum Sports.xlsx")
    wb.save(out)
    print(f"{len(want)} products -> {out}")


def load(key, xlsx):
    info, path, rows = rows_for(key)
    ws = load_workbook(xlsx, data_only=True)["Price request"]
    got = {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        sku, price, moq = r[1], r[6], r[7]
        if sku and price not in (None, ""):
            p = re.findall(r"[\d.]+", str(price))
            if p:
                got[str(sku).strip()] = (p[0], "" if moq in (None, "") else str(moq).strip())
    n = 0
    for row in rows:
        if row["sku"] in got:
            row["supplier_fob_usd"], moq = got[row["sku"]]
            if moq:
                row["supplier_moq"] = moq
            n += 1
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"{n} prices loaded into {info['sheet']}. Now run reprice.py, fetch_images.py, check photos, rebuild.")


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "make":
        make(sys.argv[2])
    elif len(sys.argv) >= 4 and sys.argv[1] == "load":
        load(sys.argv[2], sys.argv[3])
    else:
        print(__doc__)
